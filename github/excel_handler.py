import io
import os
import re
import html
from datetime import datetime
from html.parser import HTMLParser
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =========================================================================
# 1. AKILLI SAYI VE FORMAT AYRIŞTIRICI (PARSE_NUMBER)
# =========================================================================
def parse_number(val, default=0.0):
    """
    Sayısal değerleri (nokta/virgül karmaşası '105.1' vs '105,1', binlik ayracı '1.250,50' vs '1,250.50',
    birim ekleri 'kg', 'ton', 'm2', 'mm', 'mt', 'm', 'adet', 'pcs', 'tl', '$', '€' vb., parantez içi negatifler)
    güvenli ve standart bir float sayıya dönüştürür.
    """
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    
    s = str(val).strip()
    if not s or s.lower() in ('-', 'none', 'null', 'nan', 'n/a', ''):
        return default
        
    # Birim eklerini ve gereksiz karakterleri temizle
    for unit in ['kg', 'ton', 'm2', 'mm', 'mt', 'm', 'adet', 'pcs', 'tl', '$', '€']:
        if s.lower().endswith(unit):
            s = s[:len(s)-len(unit)].strip()
            
    s = s.replace(' ', '').replace('\xa0', '')
    if not s:
        return default

    # Negatif parantez kontrolü: (15.5) -> -15.5
    is_neg = False
    if s.startswith('(') and s.endswith(')'):
        is_neg = True
        s = s[1:-1].strip()

    # Hem nokta hem virgül içeriyorsa binlik ayracı tespit et
    if '.' in s and ',' in s:
        if s.rfind(',') > s.rfind('.'):
            # Avrupa / Türkiye formatı: 1.250,50 -> 1250.50
            s = s.replace('.', '').replace(',', '.')
        else:
            # Amerikan / Standart format: 1,250.50 -> 1250.50
            s = s.replace(',', '')
    elif ',' in s:
        # Tek virgül varsa ondalık ayracıdır: 105,1 -> 105.1
        s = s.replace(',', '.')
        
    try:
        val_float = float(s)
        return -val_float if is_neg else val_float
    except (ValueError, TypeError):
        return default

def safe_int(val, default=1):
    """Sayısal değeri güvenli bir pozitif tamsayıya (int) dönüştürür."""
    try:
        num = parse_number(val, float(default))
        res = int(round(num))
        return res if res > 0 else default
    except Exception:
        return default

def safe_float(val, default=0.0):
    """Geriye dönük uyumluluk takma adı (alias)"""
    return parse_number(val, default)


# =========================================================================
# 2. TEKLA HTML / TSV / XLSX TABLO OKUYUCU MOTORU
# =========================================================================
class TeklaHtmlTableParser(HTMLParser):
    """Tekla Structures tarafından üretilen HTML .xls rapor tablolarını hatasız ayrıştırır."""
    def __init__(self):
        super().__init__()
        self.tables = []
        self.current_table = None
        self.current_row = None
        self.current_cell = None
        self.in_cell = False

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == 'table':
            self.current_table = []
            self.tables.append(self.current_table)
        elif tag == 'tr':
            if self.current_table is not None:
                self.current_row = []
                self.current_table.append(self.current_row)
        elif tag in ('td', 'th'):
            if self.current_row is not None:
                self.current_cell = []
                self.in_cell = True

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == 'table':
            self.current_table = None
        elif tag == 'tr':
            self.current_row = None
        elif tag in ('td', 'th'):
            if self.current_row is not None and self.current_cell is not None:
                txt = ''.join(self.current_cell).strip()
                txt = html.unescape(txt).replace('\xa0', ' ').strip()
                self.current_row.append(txt)
                self.current_cell = None
            self.in_cell = False

    def handle_data(self, data):
        if self.in_cell and self.current_cell is not None:
            self.current_cell.append(data)

def normalize_text(s):
    """Türkçe karakterleri ve noktalama işaretlerini normalize ederek eşleştirmeyi standartlaştırır."""
    if not s:
        return ''
    s = str(s).strip().lower()
    tr_map = {'ı': 'i', 'ğ': 'g', 'ü': 'u', 'ş': 's', 'ö': 'o', 'ç': 'c', 'â': 'a', 'î': 'i'}
    for k, v in tr_map.items():
        s = s.replace(k, v)
    for ch in ['(', ')', '[', ']', ':', '.', '/', '\\', '*', '-', '_']:
        s = s.replace(ch, ' ')
    return ' '.join(s.split())

def load_table_grids(file_stream):
    """
    Verilen akıştan (file_stream, bytes, string yol) HTML tablolarını, gerçek XLSX çalışma sayfalarını
    veya CSV/TSV satırlarını otomatik tespit edip satır-sütun listeleri (grids) olarak döndürür.
    """
    if hasattr(file_stream, 'read'):
        try:
            if hasattr(file_stream, 'seek'):
                file_stream.seek(0)
            raw_bytes = file_stream.read()
            if hasattr(file_stream, 'seek'):
                file_stream.seek(0)
        except Exception:
            raw_bytes = b''
    elif isinstance(file_stream, bytes):
        raw_bytes = file_stream
    elif isinstance(file_stream, str):
        if os.path.exists(file_stream):
            with open(file_stream, 'rb') as f:
                raw_bytes = f.read()
        else:
            raw_bytes = file_stream.encode('utf-8')
    else:
        raw_bytes = b''

    grids = []
    
    # 1. Gerçek XLSX (Zip / PK başlığı) kontrolü
    if raw_bytes.startswith(b'PK\x03\x04'):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_grid = []
                for row in ws.iter_rows(values_only=True):
                    if row and any(c is not None and str(c).strip() != '' for c in row):
                        sheet_grid.append([str(c if c is not None else '').strip() for c in row])
                if sheet_grid:
                    grids.append(sheet_grid)
            return grids
        except Exception:
            pass

    # 2. Metin / HTML Çözümleme
    text_content = ''
    for enc in ['utf-8', 'utf-8-sig', 'cp1254', 'iso-8859-9', 'latin-1']:
        try:
            text_content = raw_bytes.decode(enc)
            break
        except Exception:
            continue

    if '<table' in text_content.lower() or '<tr' in text_content.lower():
        parser = TeklaHtmlTableParser()
        parser.feed(text_content)
        for t in parser.tables:
            if t and len(t) > 0:
                grids.append(t)
        return grids

    # 3. Düz Metin / CSV / TSV / Kopyala-Yapıştır Satırları
    lines = [l.strip() for l in text_content.splitlines() if l.strip()]
    if lines:
        sep = '\t' if '\t' in lines[0] else (';' if ';' in lines[0] else ',')
        grid = [[col.strip() for col in l.split(sep)] for l in lines]
        grids.append(grid)
        
    return grids

def match_columns(norm_headers):
    """Başlık satırındaki sütun isimlerini analiz edip poz, profil, ağırlık vb. indekslerini tespit eder."""
    def has_any(h, kws):
        return any(kw in h for kw in kws)

    ass_idx = None
    part_idx = None
    desc_idx = None
    prof_idx = None
    qty_idx = None
    len_idx = None
    u_wt_idx = None
    tot_wt_idx = None
    grade_idx = None

    for idx, h in enumerate(norm_headers):
        words = h.split()
        
        # 1. Toplam Ağırlık
        if has_any(h, ['toplam arlk', 'toplam agirlik', 'toplam kg', 'toplam tonaj', 'total weight', 't weight']) or ('toplam' in words and any(w in words for w in ['arlk', 'agirlik', 'kg', 'weight'])):
            if not has_any(h, ['alan', 'area', 'm2', 'm3', 'fiyat', 'tutar']):
                tot_wt_idx = idx
                continue

        # 2. Birim Ağırlık
        if has_any(h, ['birim arlk', 'birim agirlik', 'birim kg', 'unit weight', 'u weight', 'tek arlk', 'tek agirlik', 'kg m']) or ('birim' in words and any(w in words for w in ['arlk', 'agirlik', 'kg', 'weight'])):
            if not has_any(h, ['alan', 'area', 'm2', 'm3', 'fiyat', 'tutar']):
                u_wt_idx = idx
                continue

        # 3. Tanım / Açıklama / Eleman Adı
        if has_any(h, ['tanim', 'aciklama', 'description', 'parca adi', 'eleman adi', 'eleman', 'name']) or ('ad' in words and not has_any(h, ['adet', 'assembly', 'montaj', 'parca no', 'poz no', 'marka'])):
            desc_idx = idx
            continue

        # 4. Montaj Poz / Marka No
        if any(w in words for w in ['assembly', 'montaj']) or has_any(h, ['main part', 'ana poz', 'marka no', 'markasi', 'marka']):
            ass_idx = idx
            continue

        # 5. Parça Poz No
        if any(w in words for w in ['poz', 'pos', 'parca', 'part']):
            part_idx = idx
            continue

        # 6. Adet / Miktar
        if any(w in words for w in ['adet', 'miktar', 'sayi', 'qty', 'quantity', 'count']):
            qty_idx = idx
            continue

        # 7. Profil / Kesit
        if any(w in words for w in ['profil', 'kesit', 'section', 'profile']):
            prof_idx = idx
            continue

        # 8. Uzunluk / Boy
        if any(w in words for w in ['uzunluk', 'boy', 'length', 'len']):
            len_idx = idx
            continue

        # 9. Malzeme Kalitesi
        if has_any(h, ['kalite', 'grade', 'material', 'malzeme', 'celik']):
            grade_idx = idx
            continue

    return {
        'ass_idx': ass_idx,
        'part_idx': part_idx,
        'desc_idx': desc_idx,
        'prof_idx': prof_idx,
        'qty_idx': qty_idx,
        'len_idx': len_idx,
        'u_wt_idx': u_wt_idx,
        'tot_wt_idx': tot_wt_idx,
        'grade_idx': grade_idx
    }

def is_summary_row(row):
    """Tekla ve Excel raporlarının altındaki 'Toplam', 'Total', 'parts:', 'assemblies:' özet satırlarını filtreler."""
    joined = ' '.join(str(c).lower() for c in row)
    return any(k in joined for k in ['toplam', 'total', 'grand total', 'genel toplam', 'summary', 'assemblies:', 'parts:'])


# =========================================================================
# 3. EVRENSEL TEKLA STRUCTURES EXCEL OKUYUCU
# =========================================================================
def parse_tekla_excel(file_stream):
    """
    Tekla Structures Excel raporlarını (01_Part List, 02_Assembly List, 03_Assembly Part List,
    XLSX şablonları veya yapıştırılan tablolar) otomatik ayrıştırır.
    Montaj Listesi (assemblies), Montaj Parça Listesi (assembly_parts) ve Tek Parça Listesi (parts) üretir.
    """
    grids = load_table_grids(file_stream)
    
    parsed_assemblies = []
    parsed_assembly_parts = []
    parsed_parts = []

    for grid in grids:
        if len(grid) < 2:
            continue

        # Başlık satırını bul
        header_row_idx = -1
        norm_headers = []
        for r_idx, row in enumerate(grid[:15]):
            n_row = [normalize_text(c) for c in row]
            joined = ' '.join(n_row)
            if any(k in joined for k in ['assembly', 'montaj', 'poz', 'pos', 'profil', 'adet', 'qty', 'arlk', 'agirlik', 'weight']):
                header_row_idx = r_idx
                norm_headers = n_row
                break

        if header_row_idx == -1:
            continue

        cols = match_columns(norm_headers)
        ass_idx = cols['ass_idx']
        part_idx = cols['part_idx']
        desc_idx = cols['desc_idx']
        prof_idx = cols['prof_idx']
        qty_idx = cols['qty_idx']
        len_idx = cols['len_idx']
        u_wt_idx = cols['u_wt_idx']
        tot_wt_idx = cols['tot_wt_idx']
        grade_idx = cols['grade_idx']

        is_hierarchical = (ass_idx is not None and part_idx is not None and ass_idx != part_idx)
        is_pure_assembly = (ass_idx is not None and part_idx is None)
        is_pure_part = (part_idx is not None and ass_idx is None)

        current_assembly = None
        current_assembly_qty = 1
        current_assembly_desc = 'İmalat Elemanı'
        current_assembly_prof = '-'
        current_assembly_grade = 'S235JR'

        for row in grid[header_row_idx + 1:]:
            if not row or not any(str(c).strip() for c in row):
                continue
            if is_summary_row(row):
                continue

            raw_ass = row[ass_idx].strip() if ass_idx is not None and ass_idx < len(row) and row[ass_idx] else ''
            raw_part = row[part_idx].strip() if part_idx is not None and part_idx < len(row) and row[part_idx] else ''
            raw_desc = row[desc_idx].strip() if desc_idx is not None and desc_idx < len(row) and row[desc_idx] else ''
            raw_prof = row[prof_idx].strip() if prof_idx is not None and prof_idx < len(row) and row[prof_idx] else '-'
            raw_qty = row[qty_idx].strip() if qty_idx is not None and qty_idx < len(row) and row[qty_idx] else '1'
            raw_len = row[len_idx].strip() if len_idx is not None and len_idx < len(row) and row[len_idx] else '0'
            raw_uwt = row[u_wt_idx].strip() if u_wt_idx is not None and u_wt_idx < len(row) and row[u_wt_idx] else '0'
            raw_twt = row[tot_wt_idx].strip() if tot_wt_idx is not None and tot_wt_idx < len(row) and row[tot_wt_idx] else '0'
            raw_grade = row[grade_idx].strip() if grade_idx is not None and grade_idx < len(row) and row[grade_idx] else 'S235JR'

            parsed_qty = safe_int(raw_qty, 1)
            parsed_len = parse_number(raw_len, 0.0)
            parsed_uwt = parse_number(raw_uwt, 0.0)
            parsed_twt = parse_number(raw_twt, 0.0)

            # Toplam & Birim Ağırlık Uzlaştırması
            if parsed_twt == 0.0 and parsed_uwt > 0.0:
                parsed_twt = round(parsed_qty * parsed_uwt, 2)
            elif parsed_uwt == 0.0 and parsed_twt > 0.0 and parsed_qty > 0:
                parsed_uwt = round(parsed_twt / parsed_qty, 2)

            # --- DURUM 1: Hiyerarşik Montaj + Parça Raporu (03_Assembly Part List) ---
            if is_hierarchical:
                if raw_ass:
                    # Montaj Başlık Satırı
                    current_assembly = raw_ass
                    current_assembly_qty = parsed_qty
                    current_assembly_desc = raw_desc or (raw_part if raw_part and raw_part != raw_ass else 'İmalat Elemanı')
                    current_assembly_prof = raw_prof if raw_prof != '-' else '-'
                    current_assembly_grade = raw_grade or 'S235JR'

                    parsed_assemblies.append({
                        'assembly_pos': current_assembly,
                        'description': current_assembly_desc,
                        'profile_type': current_assembly_prof,
                        'quantity': current_assembly_qty,
                        'unit_weight': round(parsed_uwt, 2),
                        'total_weight': round(parsed_twt, 2),
                        'material_grade': current_assembly_grade
                    })
                elif raw_part and current_assembly:
                    # Montajın Alt Parça Satırı
                    qty_per_ass = parsed_qty
                    tot_q = qty_per_ass * current_assembly_qty
                    part_row_tot_wt = round(parsed_uwt * tot_q, 2) if parsed_uwt > 0 else round(parsed_twt * current_assembly_qty, 2)

                    parsed_assembly_parts.append({
                        'assembly_pos': current_assembly,
                        'part_pos': raw_part,
                        'description': raw_desc or 'Montaj Parçası',
                        'profile_type': raw_prof,
                        'quantity_per_assembly': qty_per_ass,
                        'total_quantity': tot_q,
                        'length': round(parsed_len, 1),
                        'unit_weight': round(parsed_uwt, 2),
                        'total_weight': round(part_row_tot_wt, 2),
                        'material_grade': raw_grade
                    })

                    parsed_parts.append({
                        'pos_no': raw_part,
                        'name': raw_desc or 'Montaj Parçası',
                        'profile_type': raw_prof,
                        'quantity': tot_q,
                        'length': round(parsed_len, 1),
                        'unit_weight': round(parsed_uwt, 2),
                        'total_weight': round(part_row_tot_wt, 2),
                        'material_grade': raw_grade
                    })

            # --- DURUM 2: Saf Montaj Listesi (02_Assembly List) ---
            elif is_pure_assembly:
                ass_code = raw_ass
                if not ass_code:
                    continue
                parsed_assemblies.append({
                    'assembly_pos': ass_code,
                    'description': raw_desc or 'İmalat Elemanı',
                    'profile_type': raw_prof,
                    'quantity': parsed_qty,
                    'unit_weight': round(parsed_uwt, 2),
                    'total_weight': round(parsed_twt, 2),
                    'material_grade': raw_grade
                })

            # --- DURUM 3: Saf Parça Listesi (01_Part List) ---
            elif is_pure_part:
                p_code = raw_part
                if not p_code:
                    continue
                parsed_parts.append({
                    'pos_no': p_code,
                    'name': raw_desc or 'Poz Parçası',
                    'profile_type': raw_prof,
                    'quantity': parsed_qty,
                    'length': round(parsed_len, 1),
                    'unit_weight': round(parsed_uwt, 2),
                    'total_weight': round(parsed_twt, 2),
                    'material_grade': raw_grade
                })

            # --- DURUM 4: Genel Tekil Satır ---
            elif raw_ass or raw_part:
                main_code = raw_ass or raw_part
                parsed_assemblies.append({
                    'assembly_pos': main_code,
                    'description': raw_desc or 'İmalat Elemanı',
                    'profile_type': raw_prof,
                    'quantity': parsed_qty,
                    'unit_weight': round(parsed_uwt, 2),
                    'total_weight': round(parsed_twt, 2),
                    'material_grade': raw_grade
                })
                parsed_parts.append({
                    'pos_no': main_code,
                    'name': raw_desc or 'Poz Parçası',
                    'profile_type': raw_prof,
                    'quantity': parsed_qty,
                    'length': round(parsed_len, 1),
                    'unit_weight': round(parsed_uwt, 2),
                    'total_weight': round(parsed_twt, 2),
                    'material_grade': raw_grade
                })

    # Tekilleştirme & Özetleme
    unique_assemblies = {}
    for a in parsed_assemblies:
        k = a['assembly_pos']
        if k not in unique_assemblies:
            unique_assemblies[k] = dict(a)
        else:
            unique_assemblies[k]['quantity'] += a['quantity']
            unique_assemblies[k]['total_weight'] = round(unique_assemblies[k]['total_weight'] + a['total_weight'], 2)

    unique_parts = {}
    for p in parsed_parts:
        k = p['pos_no']
        if k not in unique_parts:
            unique_parts[k] = dict(p)
        else:
            unique_parts[k]['quantity'] += p['quantity']
            unique_parts[k]['total_weight'] = round(unique_parts[k]['total_weight'] + p['total_weight'], 2)

    return {
        'assemblies': list(unique_assemblies.values()),
        'assembly_parts': parsed_assembly_parts,
        'parts': list(unique_parts.values())
    }

def parse_assemblies_file(file_stream):
    """Sadece Montaj / Assembly Listesi içeren Excel/HTML dosyasını ayrıştırır."""
    res = parse_tekla_excel(file_stream)
    return res['assemblies']

def parse_assembly_parts_file(file_stream):
    """Sadece Montaj Parça Listesi içeren Excel/HTML dosyasını ayrıştırır."""
    res = parse_tekla_excel(file_stream)
    if res['assembly_parts']:
        return res['assembly_parts']
    ap_list = []
    for a in res['assemblies']:
        ap_list.append({
            'assembly_pos': a['assembly_pos'],
            'part_pos': a['assembly_pos'],
            'description': a.get('description', ''),
            'profile_type': a.get('profile_type', ''),
            'quantity_per_assembly': 1,
            'total_quantity': a.get('quantity', 1),
            'length': 0.0,
            'unit_weight': a.get('unit_weight', 0.0),
            'total_weight': a.get('total_weight', 0.0),
            'material_grade': a.get('material_grade', 'S235JR')
        })
    return ap_list

def parse_parts_file(file_stream):
    """Sadece Tek Parça / Poz Listesi içeren Excel/HTML dosyasını ayrıştırır."""
    res = parse_tekla_excel(file_stream)
    return res['parts']

def parse_clipboard_table(text, target_type='assemblies'):
    """
    Excel veya Tekla'dan kopyalanan tablo metnini (Tab veya noktalı virgül / virgül ayrılmış) ayrıştırır.
    target_type: 'assemblies' | 'assembly_parts' | 'parts'
    """
    if not text or not text.strip():
        return []

    res = parse_tekla_excel(text)
    if target_type == 'assemblies':
        return res['assemblies'] or res['parts']
    elif target_type == 'assembly_parts':
        if res['assembly_parts']:
            return res['assembly_parts']
        return parse_assembly_parts_file(text)
    elif target_type == 'parts':
        return res['parts'] or res['assemblies']
    return []

def parse_excel_parts(file_stream):
    """Geriye dönük uyumluluk takma adı (alias)"""
    res = parse_tekla_excel(file_stream)
    return res['parts']


# =========================================================================
# 4. ÖRNEK EXCEL ŞABLONU OLUŞTURUCU
# =========================================================================
def generate_template_excel():
    """Kullanıcının doldurması için 2 sayfalı zengin Tekla & İmalat Excel şablonu üretir."""
    wb = openpyxl.Workbook()
    
    # 1. Sayfa: Montaj Listesi (Assemblies)
    ws1 = wb.active
    ws1.title = "1-Montaj Listesi (Assembly)"

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Koyu Mavi
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers_ass = ["Montaj / Marka No *", "Eleman Açıklaması *", "Ana Profil Kesiti *", "Adet *", "Birim Ağırlık (kg) *", "Toplam Ağırlık (kg)", "Çelik Kalitesi"]
    for c_idx, h in enumerate(headers_ass, 1):
        cell = ws1.cell(row=1, column=c_idx, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = align_center; cell.border = thin_border

    sample_assemblies = [
        ["C-101", "Ana Taşıyıcı Kolon", "HEB 300", 12, 1250.0, "=D2*E2", "S275JR"],
        ["C-102", "Rüzgar Kolonu", "HEA 240", 8, 850.0, "=D3*E3", "S275JR"],
        ["B-201", "Ana Çatı Kirişi", "IPE 360", 16, 680.0, "=D4*E4", "S355JR"],
        ["TR-301", "Çatı Makası", "2x L 80x80x8", 10, 420.0, "=D5*E5", "S235JR"],
        ["P-401", "Aşık Profili", "C 200x75x2.5", 60, 45.0, "=D6*E6", "S235JR"]
    ]

    for r_idx, r_data in enumerate(sample_assemblies, 2):
        for c_idx, val in enumerate(r_data, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10); cell.border = thin_border
            if c_idx in (1, 7): cell.alignment = align_center
            elif c_idx in (4, 5, 6): cell.alignment = align_right
            else: cell.alignment = align_left

    # 2. Sayfa: Tek Parça / Poz Listesi (Single Parts)
    ws2 = wb.create_sheet(title="2-Tek Parça Listesi (Pozlar)")
    headers_parts = ["Poz No *", "Parça Tanımı *", "Profil / Sac Kesiti *", "Adet *", "Uzunluk (mm)", "Birim Ağırlık (kg) *", "Toplam Ağırlık (kg)", "Malzeme Kalitesi"]
    for c_idx, h in enumerate(headers_parts, 1):
        cell = ws2.cell(row=1, column=c_idx, value=h)
        cell.font = header_font; cell.fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid') # Teal
        cell.alignment = align_center; cell.border = thin_border

    sample_parts = [
        ["p1", "Kolon Gövde Profili", "HEB 300", 12, 8500, 1000.0, "=D2*F2", "S275JR"],
        ["p2", "Taban Flanş Plakası", "PL 25mm", 12, 450, 45.2, "=D3*F3", "S355JR"],
        ["p3", "Berkitme Levhası", "PL 12mm", 48, 280, 5.3, "=D4*F4", "S275JR"],
        ["p4", "Bayrak Plakası", "PL 10mm", 32, 200, 3.1, "=D5*F5", "S235JR"],
        ["p5", "Kiriş Gövde", "IPE 360", 16, 12000, 685.0, "=D6*F6", "S355JR"]
    ]

    for r_idx, r_data in enumerate(sample_parts, 2):
        for c_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10); cell.border = thin_border
            if c_idx in (1, 8): cell.alignment = align_center
            elif c_idx in (4, 5, 6, 7): cell.alignment = align_right
            else: cell.alignment = align_left

    # Sütun genişliklerini otomatik ayarla
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =========================================================================
# 5. TEDARİKÇİ MALZEME TEKLİF İSTEK FORMU (RFQ EXCEL)
# =========================================================================
def export_material_rfq_excel(project_code, ordered_rows, supplier_name=""):
    """Sipariş modülündeki malzemeler için resmi teklif istek Excel formu oluşturur."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Malzeme Teklif Listesi"

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    data_font = Font(name='Segoe UI', size=10)
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    kpi_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Başlık Alanı
    ws.merge_cells('A1:F1')
    ws['A1'] = "MALZEME VE PROFİL TEKLİF İSTEK FORMU"
    ws['A1'].font = Font(name='Segoe UI', size=15, bold=True, color='1E3A8A')
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30

    ws['A2'] = "Firma:"; ws['A2'].font = bold_font; ws['B2'] = supplier_name or "-"
    ws['D2'] = "Proje Kodu:"; ws['D2'].font = bold_font; ws['E2'] = project_code
    ws['A3'] = "Tarih:"; ws['A3'].font = bold_font; ws['B3'] = datetime.now().strftime('%d.%m.%Y')
    ws['D3'] = "Konu:"; ws['D3'].font = bold_font; ws['E3'] = "Hammadde / Profil Fiyat Teklifi"
    
    table_start_row = 5
    headers = ["Malzeme / Kesit", "Kalınlık (mm)", "En (mm)", "Boy (mm)", "Adet", "Toplam Ağırlık (kg)"]
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=table_start_row, column=c_idx, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[table_start_row].height = 25

    total_weight = 0.0
    for idx, r in enumerate(ordered_rows, 1):
        row_num = table_start_row + idx
        m = r.get('material', '')
        kal = r.get('thickness', 0)
        en = r.get('width', 0)
        boy = r.get('length', 0)
        adet = r.get('quantity', 0)
        ag = r.get('weight', 0.0)
        total_weight += ag

        vals = [
            m,
            kal if kal > 0 else "-",
            en if en > 0 else "-",
            boy if boy > 0 else "-",
            adet,
            round(ag, 2)
        ]

        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            cell.font = data_font; cell.border = thin_border
            if c_idx in (2, 3, 4): cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in (5, 6): cell.alignment = Alignment(horizontal='right', vertical='center')
            else: cell.alignment = Alignment(horizontal='left', vertical='center')

    # Toplam Satırı
    tot_row = table_start_row + len(ordered_rows) + 1
    ws.cell(row=tot_row, column=4, value="GENEL TOPLAM:").font = bold_font
    ws.cell(row=tot_row, column=4).alignment = Alignment(horizontal='right', vertical='center')
    
    tot_cell = ws.cell(row=tot_row, column=6, value=round(total_weight, 2))
    tot_cell.font = bold_font; tot_cell.alignment = Alignment(horizontal='right', vertical='center')

    for c in range(1, 7):
        ws.cell(row=tot_row, column=c).fill = kpi_fill
        ws.cell(row=tot_row, column=c).border = thin_border

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =========================================================================
# 6. PROJE DETAYLI İMALAT & TONAJ İCMAL RAPORU EXCEL
# =========================================================================
def export_project_report_excel(project_summary, assemblies_list, parts_list):
    """Bir projenin Montaj, Parça ve Aşama durumlarını içeren zengin Excel raporu üretir."""
    wb = openpyxl.Workbook()
    
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    data_font = Font(name='Segoe UI', size=10)
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    kpi_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    done_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Sayfa: Montaj / Assembly Listesi
    ws1 = wb.active
    ws1.title = "Montaj Durum İcmali"

    ws1.merge_cells('A1:K1')
    ws1['A1'] = f"İMALAT & MONTAJ İCMALİ - {project_summary['name'].upper()} ({project_summary['code']})"
    ws1['A1'].font = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
    ws1['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws1.row_dimensions[1].height = 30

    info_rows = [
        ("Proje Kodu:", project_summary['code'], "Müşteri:", project_summary.get('customer') or '-'),
        ("Toplam Tonaj:", f"{project_summary['total_tonnage']} Ton", "Şantiye:", project_summary.get('site_location') or '-'),
        ("İmalatı Biten:", f"{project_summary['fab_completed_tonnage']} Ton (%{project_summary['fab_pct']})", "Boyanan:", f"{project_summary['paint_completed_tonnage']} Ton (%{project_summary['paint_pct']})"),
        ("Sevk Edilen:", f"{project_summary['shipped_tonnage']} Ton (%{project_summary['ship_pct']})", "Atölye Stoğu:", f"{project_summary['factory_stock_tonnage']} Ton")
    ]

    for r_idx, (k1, v1, k2, v2) in enumerate(info_rows, 3):
        ws1.cell(row=r_idx, column=1, value=k1).font = bold_font
        ws1.cell(row=r_idx, column=2, value=v1).font = data_font
        ws1.cell(row=r_idx, column=4, value=k2).font = bold_font
        ws1.cell(row=r_idx, column=5, value=v2).font = data_font
        for c in range(1, 6): ws1.cell(row=r_idx, column=c).fill = kpi_fill

    table_start_row = 8
    headers1 = ["Sıra", "Montaj No", "Tanım", "Profil", "Toplam Adet", "Birim Kg", "Toplam Kg", "Çatım", "Kaynak", "Boya", "Sevk"]
    for c_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=table_start_row, column=c_idx, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws1.row_dimensions[table_start_row].height = 25

    for idx, a in enumerate(assemblies_list, 1):
        r = table_start_row + idx
        vals = [
            idx,
            a['assembly_pos'],
            a['description'],
            a['profile_type'],
            a['quantity'],
            a['unit_weight'],
            a['total_weight'],
            f"{a.get('fab_fitup_qty', 0)} / {a['quantity']}",
            f"{a.get('fab_welding_qty', 0)} / {a['quantity']}",
            f"{a.get('paint_done_qty', 0)} / {a['quantity']}",
            f"{a.get('shipped_qty', 0)} / {a['quantity']}"
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws1.cell(row=r, column=c_idx, value=val)
            cell.font = data_font; cell.border = thin_border
            if c_idx in (1, 2, 8, 9, 10, 11): cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in (5, 6, 7): cell.alignment = Alignment(horizontal='right', vertical='center')
            else: cell.alignment = Alignment(horizontal='left', vertical='center')

    # 2. Sayfa: Parça Listesi (Single Parts)
    ws2 = wb.create_sheet(title="Parça Listesi & Kesim")
    headers2 = ["Sıra", "Poz No", "Parça Tanımı", "Profil Kesiti", "Toplam Adet", "Kesilen Adet", "Kalan Adet", "Uzunluk (mm)", "Birim Kg", "Toplam Kg", "Kalite"]
    for c_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=c_idx, value=h)
        cell.font = header_font; cell.fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center'); cell.border = thin_border
    ws2.row_dimensions[1].height = 25

    for idx, p in enumerate(parts_list, 1):
        r = 1 + idx
        kalan = max(0, p['quantity'] - p.get('cut_quantity', 0))
        vals = [
            idx,
            p['pos_no'],
            p['name'],
            p['profile_type'],
            p['quantity'],
            p.get('cut_quantity', 0),
            kalan,
            p.get('length', 0),
            p['unit_weight'],
            p['total_weight'],
            p['material_grade']
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c_idx, value=val)
            cell.font = data_font; cell.border = thin_border
            if c_idx in (1, 2, 11): cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in (5, 6, 7, 8, 9, 10): cell.alignment = Alignment(horizontal='right', vertical='center')
            else: cell.alignment = Alignment(horizontal='left', vertical='center')
            if c_idx == 7 and kalan == 0: cell.fill = done_fill

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =========================================================================
# 7. SEVKİYAT İRSALİYESİ / ÇEKİ LİSTESİ EXCEL ÇIKTISI
# =========================================================================
def export_shipment_excel(shipment, items, settings):
    """Sevk İrsaliyesi / Çeki Listesini resmi formatta Excel olarak oluşturur."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sevk İrsaliyesi"

    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    data_font = Font(name='Segoe UI', size=10)
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    kpi_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Başlık ve Firma Bilgisi
    ws.merge_cells('A1:F1')
    ws['A1'] = settings.get('company_name', 'ÇELİK VE METAL İMALAT SANAYİ')
    ws['A1'].font = Font(name='Segoe UI', size=13, bold=True, color='1E3A8A')
    
    ws.merge_cells('A2:F2')
    ws['A2'] = "SEVKİYAT İRSALİYESİ & ÇEKİ LİSTESİ (PACKING LIST)"
    ws['A2'].font = Font(name='Segoe UI', size=12, bold=True, color='0F172A')

    ws['A4'] = "İrsaliye No:"; ws['A4'].font = bold_font; ws['B4'] = shipment.get('dispatch_no', '-')
    ws['D4'] = "Sevk Tarihi:"; ws['D4'].font = bold_font; ws['E4'] = shipment.get('dispatch_date', '-')
    
    ws['A5'] = "Araç Plakası:"; ws['A5'].font = bold_font; ws['B5'] = shipment.get('vehicle_plate', '-')
    ws['D5'] = "Şoför / İletişim:"; ws['D5'].font = bold_font; ws['E5'] = f"{shipment.get('driver_name', '-')} ({shipment.get('driver_phone', '-')})"
    
    ws['A6'] = "Teslim Yeri:"; ws['A6'].font = bold_font; ws['B6'] = shipment.get('destination', '-')
    ws['D6'] = "Nakliye Firması:"; ws['D6'].font = bold_font; ws['E6'] = shipment.get('carrier_company', '-')

    table_start_row = 8
    headers = ["Sıra", "Montaj / Marka No", "Açıklama / Eleman", "Sevk Edilen Adet", "Birim Ağırlık (kg)", "Toplam Ağırlık (kg)"]
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=table_start_row, column=c_idx, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[table_start_row].height = 25

    total_qty = 0
    total_weight = 0.0

    for idx, it in enumerate(items, 1):
        r = table_start_row + idx
        qty = it.get('quantity', 1)
        u_wt = it.get('unit_weight', 0.0)
        t_wt = it.get('total_weight', qty * u_wt)
        total_qty += qty
        total_weight += t_wt

        vals = [idx, it.get('assembly_pos', '-'), it.get('description', '-'), qty, round(u_wt, 2), round(t_wt, 2)]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c_idx, value=val)
            cell.font = data_font; cell.border = thin_border
            if c_idx in (1, 2): cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in (4, 5, 6): cell.alignment = Alignment(horizontal='right', vertical='center')
            else: cell.alignment = Alignment(horizontal='left', vertical='center')

    # Toplam Satırı
    tot_row = table_start_row + len(items) + 1
    ws.cell(row=tot_row, column=3, value="GENEL TOPLAM:").font = bold_font
    ws.cell(row=tot_row, column=3).alignment = Alignment(horizontal='right', vertical='center')
    
    ws.cell(row=tot_row, column=4, value=total_qty).font = bold_font
    ws.cell(row=tot_row, column=4).alignment = Alignment(horizontal='right', vertical='center')

    ws.cell(row=tot_row, column=6, value=round(total_weight, 2)).font = bold_font
    ws.cell(row=tot_row, column=6).alignment = Alignment(horizontal='right', vertical='center')

    for c in range(1, 7):
        ws.cell(row=tot_row, column=c).fill = kpi_fill
        ws.cell(row=tot_row, column=c).border = thin_border

    # İmza Alanları
    sign_row = tot_row + 3
    ws.cell(row=sign_row, column=1, value="Teslim Eden (Fabrika)").font = bold_font
    ws.cell(row=sign_row, column=3, value="Taşıyan (Şoför)").font = bold_font
    ws.cell(row=sign_row, column=5, value="Teslim Alan (Şantiye)").font = bold_font

    ws.cell(row=sign_row + 1, column=1, value="İmza: ..................")
    ws.cell(row=sign_row + 1, column=3, value="İmza: ..................")
    ws.cell(row=sign_row + 1, column=5, value="İmza: ..................")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
